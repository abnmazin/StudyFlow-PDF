import telebot
from telebot import types
from openpyxl import Workbook, load_workbook
import os
import threading
import json
from datetime import datetime
import traceback

# ── الإعدادات الأساسية ──────────────────────────────────────────────────
ADMIN_IDS = [1360000077, 6382645701,7238142091] 
#BOT_TOKEN = '1644721136:AAGswGqerNZntSvp9-yEZKF_SMjg1PgQ8VM'  
BOT_TOKEN = '8629204913:AAFP5IZN9JC5BHXPv-ljm5BZ-NnQvl51OtM'  

FILE_STAGE_2 = 'stage_2.xlsx'
FILE_STAGE_3 = 'stage_3.xlsx'
SESSIONS_FILE = 'sessions.json'

COMPANIES_DATA = [
    {"name": "مركز التدريب المهني", "s2m7": 0, "s2m8": 1, "s3m7": 1, "s3m8": 0}, # الإجمالي 2 (1+1)
    {"name": "شركة الحفر العراقية", "s2m7": 2, "s2m8": 2, "s3m7": 3, "s3m8": 3}, # الإجمالي 10
    {"name": "المديرية العامة لإنتاج الطاقة", "s2m7": 15, "s2m8": 15, "s3m7": 23, "s3m8": 22}, # الإجمالي 75
    {"name": "المديرية العامة لنقل الطاقة", "s2m7": 2, "s2m8": 2, "s3m7": 3, "s3m8": 3}, # الإجمالي 10
    {"name": "توزيع كهرباء الجنوب (البصرة)", "s2m7": 4, "s2m8": 4, "s3m7": 6, "s3m8": 6}, # الإجمالي 20
    {"name": "توزيع كهرباء الجنوب (شمال البصرة)", "s2m7": 4, "s2m8": 4, "s3m7": 6, "s3m8": 6}, # الإجمالي 20
    {"name": "الشركة العامة لموانئ العراق", "s2m7": 1, "s2m8": 2, "s3m7": 2, "s3m8": 2}, # الإجمالي 7
    {"name": "الشركة العامة للحديد والصلب", "s2m7": 2, "s2m8": 2, "s3m7": 3, "s3m8": 3}, # الإجمالي 10
    {"name": "الشركة العامة لصناعة الاسمدة", "s2m7": 2, "s2m8": 2, "s3m7": 3, "s3m8": 3}, # الإجمالي 10
    {"name": "شركة توزيع المنتجات النفطية", "s2m7": 1, "s2m8": 2, "s3m7": 2, "s3m8": 2}, # الإجمالي 7
    {"name": "الشركة العامة للبتروكيمياويات", "s2m7": 2, "s2m8": 2, "s3m7": 3, "s3m8": 3}, # الإجمالي 10
    {"name": "مديرية بريد واتصالات", "s2m7": 4, "s2m8": 4, "s3m7": 6, "s3m8": 6}, # الإجمالي 20
    {"name": "شركة تعبئة وخدمات الغاز", "s2m7": 1, "s2m8": 1, "s3m7": 1, "s3m8": 2}, # الإجمالي 5 (3+2)
    {"name": "معهد التدريب النفطي", "s2m7": 0, "s2m8": 1, "s3m7": 1, "s3m8": 0}, # الإجمالي 2
    {"name": "شركة ابن ماجد", "s2m7": 4, "s2m8": 4, "s3m7": 6, "s3m8": 6}, # الإجمالي 20
]

bot = telebot.TeleBot(BOT_TOKEN)

# ── معالج Callback Query لمنع تعليق الأزرار ─────────────────────────────
@bot.callback_query_handler(func=lambda c: True)
def handle_all_callbacks(call):
    # كل الأزرار حالياً للعرض فقط (callback_data = "ignore")
    # نرسل إشعاراً سريعاً يخفي مؤشر التحميل فوراً دون تغيير أي شيء
    bot.answer_callback_query(call.id, text="", show_alert=False)

# أقفال الحماية (منفصلة لكل مرحلة لضمان السرعة القصوى)
lock_stage_2 = threading.RLock()
lock_stage_3 = threading.RLock()
session_lock = threading.RLock()

def get_lock(stage):
    return lock_stage_2 if str(stage) == '2' else lock_stage_3

# حالات الجلسة
ST_ADMIN   = 'admin_home'
ST_STAGE   = 'waiting_stage'
ST_COMPANY = 'waiting_company'
ST_MONTH   = 'waiting_month'
ST_INFO    = 'waiting_info'
ST_WITHDRAW_CONFIRM = 'waiting_withdraw_confirm'

# ── تهيئة ملفات الإكسل ──────────────────────────────────────────────────
def get_file(stage):
    return FILE_STAGE_2 if str(stage) == '2' else FILE_STAGE_3

for stage, file in [('2', FILE_STAGE_2), ('3', FILE_STAGE_3)]:
    with get_lock(stage):
        if not os.path.exists(file):
            wb = Workbook()
            ws = wb.active
            ws.title = "Registrations"
            # تمت إضافة عمود Timestamp
            ws.append(['ID', 'Username', 'Name', 'Address', 'Company', 'Month', 'Timestamp', 'Status'])
            wb.save(file)

# ── دوال إدارة الجلسات (JSON) ───────────────────────────────────────────
def load_sessions():
    with session_lock:
        if not os.path.exists(SESSIONS_FILE):
            return {}
        try:
            with open(SESSIONS_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except:
            return {}

def save_sessions(sessions_data):
    with session_lock:
        with open(SESSIONS_FILE, 'w', encoding='utf-8') as f:
            json.dump(sessions_data, f, ensure_ascii=False, indent=4)

def get_session(user_id):
    return load_sessions().get(str(user_id))

def update_session(user_id, data):
    sessions = load_sessions()
    uid_str = str(user_id)
    if uid_str not in sessions:
        sessions[uid_str] = {}
    sessions[uid_str].update(data)
    save_sessions(sessions)

def clear_session(user_id):
    sessions = load_sessions()
    uid_str = str(user_id)
    if uid_str in sessions:
        del sessions[uid_str]
        save_sessions(sessions)

# ── دوال مساعدة للإكسل ──────────────────────────────────────────────────
def get_all_registrations(stage):
    with get_lock(stage):
        file = get_file(stage)
        wb = load_workbook(file)
        ws = wb.active
        rows = list(ws.values)
        if len(rows) <= 1:
            wb.close()
            return []
        headers = rows[0]
        result = [dict(zip(headers, row)) for row in rows[1:]]
        wb.close()
        return result

def is_registered_in_any(user_id):
    for stage in ['2', '3']:
        regs = get_all_registrations(stage)
        if any(str(r['ID']) == str(user_id) and r['Status'] == 'confirmed' for r in regs):
            return True
    return False

def delete_registration(user_id):
    for stage in ['2', '3']:
        file = get_file(stage)
        with get_lock(stage):
            if os.path.exists(file):
                wb = load_workbook(file)
                ws = wb.active
                for row in range(ws.max_row, 1, -1):
                    if str(ws.cell(row=row, column=1).value) == str(user_id):
                        ws.delete_rows(row)
                wb.save(file)
                wb.close()

def get_capacity(company_name, stage, month):
    for c in COMPANIES_DATA:
        if c['name'] == company_name:
            return c[f"s{stage}m{month}"]
    return 0

def available_slots(company_name, stage, month, regs):
    cap = get_capacity(company_name, stage, month)
    taken = sum(1 for r in regs if r['Company'] == company_name and str(r['Month']) == str(month) and r['Status'] == 'confirmed')
    return max(0, cap - taken)

def get_avail_companies(stage, regs):
    avail = []
    for c in COMPANIES_DATA:
        m7 = available_slots(c['name'], stage, '7', regs)
        m8 = available_slots(c['name'], stage, '8', regs)
        if m7 > 0 or m8 > 0:
            avail.append(c['name'])
    return avail

# ── أزرار الكيبورد ──────────────────────────────────────────────────────
def admin_kb():
    kb = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
    kb.add("📊 عرض الإحصائيات", "📥 تنزيل الملفات")
    kb.add("🎓 الدخول كطالب")
    return kb

def stage_kb():
    kb = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
    kb.add("المرحلة الثانية", "المرحلة الثالثة")
    return kb

def companies_kb(available_names):
    kb = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
    for name in available_names:
        kb.add(name)
    kb.add("رجوع")
    return kb

def months_kb(slots_7, slots_8):
    kb = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=2)
    if slots_7 > 0:
        kb.add(f"شهر 7 ({slots_7} مقاعد)")
    if slots_8 > 0:
        kb.add(f"شهر 8 ({slots_8} مقاعد)")
    kb.add("رجوع")
    return kb

def withdraw_kb():
    kb = types.ReplyKeyboardMarkup(resize_keyboard=True, row_width=1)
    kb.add("سحب التسجيل وإعادة الاختيار", "إلغاء العملية")
    return kb

# ── دوال الآدمن المستقلة ────────────────────────────────────────────────
def cmd_admin_stats(chat_id):
    for stage in ['2', '3']:
        regs = get_all_registrations(stage)
        confirmed = [r for r in regs if r['Status'] == 'confirmed']
        
        # إنشاء الكيبورد الشفاف لهذه المرحلة
        kb = types.InlineKeyboardMarkup(row_width=2)
        
        for c in COMPANIES_DATA:
            name = c['name']
            cap7 = get_capacity(name, stage, '7')
            cap8 = get_capacity(name, stage, '8')
            
            # إذا كانت الشركة لا تدعم هذه المرحلة، نتخطاها
            if cap7 == 0 and cap8 == 0:
                continue

            taken7 = sum(1 for r in confirmed if r['Company'] == name and str(r['Month']) == '7')
            taken8 = sum(1 for r in confirmed if r['Company'] == name and str(r['Month']) == '8')
            
            # 1. زر اسم الشركة (عريض يأخذ الصف كاملاً)
            kb.add(types.InlineKeyboardButton(text=f"🏢 {name}", callback_data="ignore"))
            
            # 2. أزرار الشهور (زرين في صف واحد تحت اسم الشركة)
            row_buttons = []
            if cap7 > 0:
                status7 = "🔴" if taken7 >= cap7 else "🟢"
                row_buttons.append(types.InlineKeyboardButton(
                    text=f"ش7: {taken7}/{cap7} {status7}", 
                    callback_data="ignore"
                ))
            if cap8 > 0:
                status8 = "🔴" if taken8 >= cap8 else "🟢"
                row_buttons.append(types.InlineKeyboardButton(
                    text=f"ش8: {taken8}/{cap8} {status8}", 
                    callback_data="ignore"
                ))
            
            # إضافة صف أزرار الشهور للكيبورد
            if row_buttons:
                kb.row(*row_buttons)

        stage_title = (
            f"📊 *إحصائيات المرحلة {'الثانية' if stage=='2' else 'الثالثة'}*\n"
            f"──── ──── ──── ────\n"
            f"✅ إجمالي المسجلين: {len(confirmed)}"
        )
        
        bot.send_message(chat_id, stage_title, reply_markup=kb, parse_mode='Markdown')

def cmd_download_files(chat_id):
    bot.send_message(chat_id, "⏳ جاري تجهيز الملفات...")
    for stage, file_path in [('2', FILE_STAGE_2), ('3', FILE_STAGE_3)]:
        if os.path.exists(file_path):
            # نقرأ الملف بسرعة داخل القفل كـ Bytes حتى لا نعرقل الطلاب
            with get_lock(stage):
                with open(file_path, 'rb') as f:
                    file_data = f.read()
            
            # نرسل الملف خارج القفل
            bot.send_document(chat_id, (file_path, file_data), caption=f"ملف المرحلة {'الثانية' if stage=='2' else 'الثالثة'}")

# ── معالج الرسائل الرئيسي ───────────────────────────────────────────────
@bot.message_handler(func=lambda m: True)
def handle_all_messages(msg):
    uid = msg.from_user.id
    text = msg.text.strip()

    if text == '/start':
        if uid in ADMIN_IDS:
            update_session(uid, {'state': ST_ADMIN})
            bot.send_message(msg.chat.id, "أهلاً بك في لوحة الإدارة. اختر الإجراء المطلوب:", reply_markup=admin_kb())
            return
            
        if is_registered_in_any(uid):
            update_session(uid, {'state': ST_WITHDRAW_CONFIRM})
            warning_msg = (
                "أنت مسجل مسبقاً في النظام.\n\n"
                "⚠️ *تحذير هام:*\n"
                "هل ترغب في سحب تسجيلك الحالي وإعادة الاختيار؟\n"
                "إجراء السحب نهائي ولا يمكن التراجع عنه، وقد تفقد مقعدك الحالي لصالح طالب آخر."
            )
            bot.send_message(msg.chat.id, warning_msg, reply_markup=withdraw_kb(), parse_mode='Markdown')
            return
            
        clear_session(uid)
        update_session(uid, {'state': ST_STAGE})
        bot.send_message(
            msg.chat.id,
            "مرحباً بك في نظام تسجيل التدريب الصيفي.\nيرجى اختيار المرحلة الدراسية:",
            reply_markup=stage_kb()
        )
        return

    sess = get_session(uid)
    if not sess:
        bot.send_message(msg.chat.id, "انتهت الجلسة. يرجى إرسال /start للبدء من جديد.", reply_markup=types.ReplyKeyboardRemove())
        return

    state = sess['state']

    # ── لوحة الآدمن ──
    if state == ST_ADMIN and uid in ADMIN_IDS:
        if text == "📊 عرض الإحصائيات":
            cmd_admin_stats(msg.chat.id)
        elif text == "📥 تنزيل الملفات":
            cmd_download_files(msg.chat.id)
        elif text == "🎓 الدخول كطالب":
            clear_session(uid)
            if is_registered_in_any(uid):
                update_session(uid, {'state': ST_WITHDRAW_CONFIRM})
                warning_msg = (
                    "أنت مسجل مسبقاً في النظام.\n\n"
                    "⚠️ *تحذير هام:*\n"
                    "هل ترغب في سحب تسجيلك الحالي والتسجيل كطالب؟\n"
                    "إجراء السحب نهائي ولا يمكن التراجع عنه، وقد يفقد مقعدك الحالي لصالح طالب آخر."
                )
                bot.send_message(msg.chat.id, warning_msg, reply_markup=withdraw_kb(), parse_mode='Markdown')
                return
            update_session(uid, {'state': ST_STAGE})
            bot.send_message(msg.chat.id, "تم تحويلك لواجهة الطالب.\nيرجى اختيار المرحلة الدراسية:", reply_markup=stage_kb())
        else:
            bot.send_message(msg.chat.id, "يرجى اختيار أمر من القائمة.")
        return

    # ── معالجة تأكيد السحب ──
    if state == ST_WITHDRAW_CONFIRM:
        if text == "سحب التسجيل وإعادة الاختيار":
            delete_registration(uid)
            clear_session(uid)
            update_session(uid, {'state': ST_STAGE})
            bot.send_message(
                msg.chat.id, 
                "تم سحب تسجيلك وإلغاء حجز مقعدك بنجاح. يمكنك الآن إعادة التسجيل.\n\nيرجى اختيار المرحلة الدراسية:", 
                reply_markup=stage_kb()
            )
        elif text == "إلغاء العملية":
            clear_session(uid)
            bot.send_message(
                msg.chat.id, 
                "تم إلغاء عملية السحب. لا يزال تسجيلك ومقعدك الحالي فعالاً.", 
                reply_markup=types.ReplyKeyboardRemove()
            )
            # إعادة الآدمن للوحته إذا ألغى السحب
            if uid in ADMIN_IDS:
                update_session(uid, {'state': ST_ADMIN})
                bot.send_message(msg.chat.id, "العودة للوحة الإدارة:", reply_markup=admin_kb())
        else:
            bot.send_message(msg.chat.id, "يرجى اختيار أحد الخيارات من القائمة.")
        return

    # ── باقي خطوات التسجيل للطلاب ──
    if state == ST_STAGE:
        if text in ["المرحلة الثانية", "المرحلة الثالثة"]:
            stage = '2' if 'الثانية' in text else '3'
            regs = get_all_registrations(stage)
            avail = get_avail_companies(stage, regs)

            if not avail:
                bot.send_message(msg.chat.id, "عذراً، لا توجد مقاعد متاحة لهذه المرحلة حالياً.", reply_markup=stage_kb())
                return

            update_session(uid, {'state': ST_COMPANY, 'stage': stage})
            bot.send_message(msg.chat.id, "يرجى اختيار جهة التدريب (الشركة):", reply_markup=companies_kb(avail))
        else:
            bot.send_message(msg.chat.id, "يرجى اختيار المرحلة من الخيارات المتاحة.")
        return

    if state == ST_COMPANY:
        if text == "رجوع":
            update_session(uid, {'state': ST_STAGE})
            bot.send_message(msg.chat.id, "يرجى اختيار المرحلة الدراسية:", reply_markup=stage_kb())
            return

        selected_company = None
        for c in COMPANIES_DATA:
            if c['name'] == text:
                selected_company = c['name']
                break

        if selected_company:
            stage = sess['stage']
            regs = get_all_registrations(stage)
            m7 = available_slots(selected_company, stage, '7', regs)
            m8 = available_slots(selected_company, stage, '8', regs)

            if m7 == 0 and m8 == 0:
                bot.send_message(msg.chat.id, "اكتملت المقاعد في هذه الجهة. يرجى اختيار جهة أخرى.")
                avail = get_avail_companies(stage, regs)
                bot.send_message(msg.chat.id, "الجهات المتاحة:", reply_markup=companies_kb(avail))
                return

            update_session(uid, {'state': ST_MONTH, 'company': selected_company})
            bot.send_message(msg.chat.id, "يرجى اختيار شهر التدريب:", reply_markup=months_kb(m7, m8))
        else:
            bot.send_message(msg.chat.id, "يرجى اختيار جهة من الخيارات المتاحة.")
        return

    if state == ST_MONTH:
        if text == "رجوع":
            update_session(uid, {'state': ST_COMPANY})
            stage = sess['stage']
            regs = get_all_registrations(stage)
            avail = get_avail_companies(stage, regs)
            bot.send_message(msg.chat.id, "يرجى اختيار جهة التدريب (الشركة):", reply_markup=companies_kb(avail))
            return

        if "شهر 7" in text or "شهر 8" in text:
            month = '7' if 'شهر 7' in text else '8'
            stage = sess['stage']
            company = sess['company']
            
            regs = get_all_registrations(stage)
            slots = available_slots(company, stage, month, regs)

            if slots <= 0:
                bot.send_message(msg.chat.id, "عذراً، اكتملت مقاعد هذا الشهر. يرجى اختيار شهر أو جهة أخرى.")
                return

            update_session(uid, {'state': ST_INFO, 'month': month})
            
            msg_text = (
                "يرجى إرسال بياناتك (الاسم الثلاثي وعنوان السكن) في رسالة واحدة، "
                "بحيث يكون الاسم في السطر الأول والعنوان في السطر الثاني.\n\n"
                "مثال:\n"
                "علي مجتبى حسن\n"
                "دور الصحة"
            )
            bot.send_message(msg.chat.id, msg_text, reply_markup=types.ReplyKeyboardRemove())
        else:
            bot.send_message(msg.chat.id, "يرجى اختيار الشهر من الخيارات المتاحة.")
        return

    if state == ST_INFO:
        lines = text.split('\n')
        
        if len(lines) < 2 or len(lines[0].split()) < 2:
            bot.send_message(
                msg.chat.id, 
                "إدخال غير صحيح. يرجى كتابة الاسم الثلاثي في السطر الأول، وعنوان السكن في السطر الثاني."
            )
            return

        name = lines[0].strip()
        address = lines[1].strip()
        
        stage = sess['stage']
        company = sess['company']
        month = sess['month']
        username = f"@{msg.from_user.username}" if msg.from_user.username else "لا يوجد"
        # توليد الطابع الزمني بأجزاء الثانية
        timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S.%f')[:-3]
        file_to_save = get_file(stage)

        with get_lock(stage):
            try:
                if is_registered_in_any(uid):
                    bot.send_message(msg.chat.id, "تم رفض الطلب: أنت مسجل مسبقاً في النظام.")
                    clear_session(uid)
                    return
                    
                regs = get_all_registrations(stage)
                if available_slots(company, stage, month, regs) <= 0:
                    bot.send_message(msg.chat.id, "عذراً، اكتملت المقاعد المتاحة أثناء محاولتك للتسجيل. يرجى إرسال /start والمحاولة مرة أخرى.")
                    clear_session(uid)
                    return

                wb = load_workbook(file_to_save)
                ws = wb.active
                # حفظ البيانات مع الطابع الزمني
                ws.append([uid, username, name, address, company, month, timestamp, 'confirmed'])
                wb.save(file_to_save)
                wb.close()

                clear_session(uid)
                success_msg = (
                    f"تم تأكيد التسجيل بنجاح.\n\n"
                    f"الاسم: {name}\n"
                    f"السكن: {address}\n"
                    f"المرحلة: {'الثانية' if stage=='2' else 'الثالثة'}\n"
                    f"الجهة: {company}\n"
                    f"الشهر: {month}\n\n"
                    'في حال هناك خطا في الاسم او العنوان قم بمراسلة ممثل المرحلة واذا كان هناك خطا في اختيار مكان التدريب ارسل /start لاعادة الاختيار'
                )
                
                # إذا كان آدمن ويسجل كطالب، نعيده للوحته
                if uid in ADMIN_IDS:
                    update_session(uid, {'state': ST_ADMIN})
                    bot.send_message(msg.chat.id, success_msg + "\n\nالعودة للوحة الإدارة:", reply_markup=admin_kb())
                else:
                    bot.send_message(msg.chat.id, success_msg)

            except PermissionError:
                bot.send_message(msg.chat.id, "النظام قيد التحديث من قبل الإدارة. يرجى المحاولة بعد قليل.")
            except Exception as e:
                bot.send_message(msg.chat.id, "حدث خطأ داخلي أثناء حفظ البيانات. يرجى إبلاغ الإدارة.")
                print(f"Error saving to Excel: {e}")
                traceback.print_exc()

if __name__ == '__main__':
    print("System is running with enhanced Admin features and precise Timestamps...")
    import time
    while True:
        try:
            bot.infinity_polling()
        except Exception as e:
            print(f"⚠️ Connection lost: {e}")
            print("🔄 Retrying in 5 seconds...")
            time.sleep(5)
            continue
        break
