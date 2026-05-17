import openpyxl, json, sys

# Force UTF-8 for stdout
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('stage_3_sorted.xlsx')
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
sorted_data = []
for r in range(2, ws.max_row + 1):
    row = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(r, c).value
        row[headers[c-1]] = str(val) if val is not None else ''
    sorted_data.append(row)

# Also read original
wb_orig = openpyxl.load_workbook('stage_3.xlsx')
ws_orig = wb_orig.active
orig_data = []
for r in range(2, ws_orig.max_row + 1):
    row = {}
    for c in range(1, ws_orig.max_column + 1):
        val = ws_orig.cell(r, c).value
        row[headers[c-1]] = str(val) if val is not None else ''
    orig_data.append(row)

print(f"Original rows: {len(orig_data)}, Sorted rows: {len(sorted_data)}")
print(f"Same count? {len(orig_data) == len(sorted_data)}")

# Verify sort order
errors = []
for i in range(1, len(sorted_data)):
    prev, curr = sorted_data[i-1], sorted_data[i]
    if prev['Company'] > curr['Company']:
        errors.append(f"Row {i+1}: Company order wrong: {prev['Company']} > {curr['Company']}")
    elif prev['Company'] == curr['Company']:
        if prev['Month'] > curr['Month']:
            errors.append(f"Row {i+1}: Month order wrong in {prev['Company']}: {prev['Month']} > {curr['Month']}")
        elif prev['Month'] == curr['Month'] and prev['Name'] > curr['Name']:
            errors.append(f"Row {i+1}: Name order wrong in {prev['Company']}/{prev['Month']}: {prev['Name']} > {curr['Name']}")

if errors:
    print(f"\n=== {len(errors)} ERRORS FOUND ===")
    for e in errors[:10]:
        print(e)
else:
    print("\nALL SORT CHECKS PASSED! No ordering errors.")

# Write results to JSON file for inspection
with open('stage_3_verify.json', 'w', encoding='utf-8') as f:
    json.dump({
        "verified": len(errors)==0,
        "errors": errors[:20],
        "first_30_rows": sorted_data[:30]
    }, f, ensure_ascii=False, indent=2)

print(f"Verification file written: stage_3_verify.json")