import openpyxl
import json

# Read the Excel file
wb = openpyxl.load_workbook('stage_3.xlsx')
ws = wb.active

# Extract headers and data
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
data = []
for r in range(2, ws.max_row + 1):
    row = {}
    for c in range(1, ws.max_column + 1):
        val = ws.cell(r, c).value
        row[headers[c-1]] = str(val) if val is not None else ''
    data.append(row)

# Sort: Company (ascending Arabic/alphabetic), Month (7 before 8), Name (ascending)
data.sort(key=lambda x: (x['Company'], x['Month'], x['Name']))

# Create a new workbook
wb_out = openpyxl.Workbook()
ws_out = wb_out.active
ws_out.title = "Registrations"

# Write headers
for c, h in enumerate(headers, 1):
    ws_out.cell(1, c, h)

# Write sorted data
for i, row in enumerate(data, 2):
    for c, h in enumerate(headers, 1):
        ws_out.cell(i, c, row[h])

# Save
wb_out.save('stage_3_sorted.xlsx')
print(f"Done! Sorted {len(data)} rows. Output: stage_3_sorted.xlsx")